"""
回收第 1 轮后，生成第 2 轮的表。

单人:
    python fill_round2.py 张医生_第1轮.xlsx A 张医生_第2轮.xlsx

批量（推荐）——把所有第 1 轮文件放进一个文件夹，自动按顺序分配 A/B/C/D:
    python fill_round2.py --batch round1_收回/ round2_待发/

输出两个文件，用途不同：
  *.xlsx            → 第 2 轮的表，发给医生
  *_分配记录.csv    → 研究者留存，分析时必需（含空白条件下的既往值，用于算落差）
"""
import sys, os, glob, shutil, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

WORD = {"C1_无既往":  None,
        "C2_本人既往": "您本人在基线时的评估：24 个月内进展为痴呆的概率 {v}%",
        "C3_转诊医生": "转诊医生在基线时的评估：24 个月内进展为痴呆的概率 {v}%",
        "C4_AI系统":  "某临床 AI 系统在基线时的输出：24 个月内进展为痴呆的概率 {v}%"}
TEMPLATE, CASES = "第2轮_完整评估_模板.xlsx", "cases.csv"

def read_round1(path):
    """读第 1 轮：G 列=风险判断，H 列=把握度。第 5 行是示例行，从第 6 行开始。"""
    ws = load_workbook(path, data_only=True)["第1轮_基线评估"]
    prior, conf, name = {}, {}, None
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row or not isinstance(row[0], str): continue
        if row[0].startswith("评估者姓名"): name = row[1]; continue
        if row[0].startswith("C") and row[0][1:].isdigit() and row[6] is not None:
            prior[row[0]] = int(row[6]); conf[row[0]] = row[7]
    return prior, conf, name

def make(r1_path, rot, out_path, quiet=False):
    rot = rot.upper(); assert rot in "ABCD", "方案必须是 A/B/C/D"
    D = pd.read_csv(CASES)
    prior, conf, name = read_round1(r1_path)
    miss = [c for c in D.case if c not in prior]
    if miss and not quiet:
        print(f"  ⚠ 第 1 轮缺 {len(miss)} 例未填: {miss[:6]}{'…' if len(miss)>6 else ''}")
        print("    这些病例第 2 轮按『无既往』处理，分析时需整例剔除。")
    bad = [c for c, v in conf.items() if v and v not in
           ("很不确定", "不太确定", "比较确定", "很确定")]
    if bad and not quiet: print(f"  ⚠ 把握度取值异常 {len(bad)} 例: {bad[:5]}")

    shutil.copy(TEMPLATE, out_path)
    wb = load_workbook(out_path); ws = wb["第2轮_完整评估"]
    ws.cell(row=1, column=1).value = f"第 2 轮：完整随访资料评估　（分配方案 {rot}）"
    n = 0
    for i, (_, r) in enumerate(D.iterrows()):
        cell = ws.cell(row=5 + i, column=6)
        tpl = WORD[r[f"rot_{rot}"]]
        if tpl and r.case in prior:
            cell.value = tpl.format(v=prior[r.case]); cell.font = Font(name="Arial", size=8.5); n += 1
        else:
            cell.value = ""
    wb.save(out_path)
    pd.DataFrame({"case": D.case, "cond": D[f"rot_{rot}"],
                  "prior_value": [prior.get(c) for c in D.case],
                  "round1_conf": [conf.get(c) for c in D.case],
                  "rotation": rot, "source_file": os.path.basename(r1_path),
                  "clinician": name}).to_csv(out_path.replace(".xlsx", "_分配记录.csv"),
                                             index=False, encoding="utf-8-sig")
    print(f"  ✓ {os.path.basename(out_path)}　方案 {rot}　"
          f"{n} 例带既往评估 / {len(D)-n} 例空白对照　评估者: {name or '未填写'}")
    return n

def batch(in_dir, out_dir):
    files = sorted(glob.glob(os.path.join(in_dir, "*.xlsx")))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files: sys.exit(f"{in_dir} 下没有找到 xlsx 文件")
    os.makedirs(out_dir, exist_ok=True)
    print(f"找到 {len(files)} 份第 1 轮回收表，按顺序分配 A/B/C/D：\n")
    for k, f in enumerate(files):
        rot = "ABCD"[k % 4]
        stem = os.path.splitext(os.path.basename(f))[0].replace("第1轮", "").strip("_ ")
        make(f, rot, os.path.join(out_dir, f"{stem}_第2轮.xlsx"))
    if len(files) % 4:
        print(f"\n  提示：{len(files)} 位医生不是 4 的倍数，四个条件的病例数会不平衡。"
              f"\n        建议凑到 {(len(files)//4+1)*4} 位，或在分析时用混合效应模型（本就推荐）。")
    print(f"\n全部生成完毕 → {out_dir}")
    print("请把 xlsx 发给对应医生；_分配记录.csv 自己留存。")

if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "--batch":
        if len(sys.argv) < 4: print(__doc__); sys.exit(1)
        batch(sys.argv[2], sys.argv[3])
    elif len(sys.argv) >= 4:
        make(sys.argv[1], sys.argv[2], sys.argv[3])
    else:
        print(__doc__); sys.exit(1)
